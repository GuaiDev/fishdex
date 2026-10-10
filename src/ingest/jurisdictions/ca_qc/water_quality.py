"""Quebec river water quality — MELCCFP "Suivi de la qualité de l'eau du fleuve et des rivières".

Source: Données Québec open data portal, dataset
``suivi-physicochimique-des-rivieres-et-du-fleuve`` (CC-BY 4.0), published by the
Ministère de l'Environnement, de la Lutte contre les changements climatiques, de
la Faune et des Parcs (MELCCFP). It is the open-data face of the Réseau-rivières
(formerly RSQER) and the IQBP index.

An earlier version of this module was a stub on the belief that RSQER was PDF-only.
That was wrong: the dataset publishes machine-readable files. Two of them matter.

STATIONS — ``IQBP_csv.zip`` (resource format CSV), member ``stations_p.csv``.
  One row per station per year (semicolon-delimited, UTF-8 with BOM) with
  coordinates and per-period summary statistics. The summaries are annual
  aggregates (IQBP index, ammonia, chlorophyll, ...) and carry no pH, temperature,
  conductivity or dissolved oxygen, so they are not stored as readings. The row
  also names the per-watershed *chiffrier* (spreadsheet) that holds the samples.

SAMPLES — the ``URL_ZGIEBV`` chiffrier named on each station row
  An ``.xlsx`` workbook per watershed per three-year period, hosted by the
  same publisher. Its ``Données transposées`` sheet has one row per sample and one
  column per parameter, including ``PH``, ``TEMP`` and ``COND``. That is where
  the readings come from.

What this does NOT provide:
  * Dissolved oxygen — Réseau-rivières does not measure it (``DBO5`` is biochemical
    oxygen demand, a different quantity), so ``do_mgl`` is always None. The
    chemistry slice reports that as "no records", which is true.
  * Turbidity — published in UTN (≈ NTU), the table column is FNU. Different
    optics, so it is left None rather than relabelled.

Values are stored as published. Station names are the French ``DESCRIPTION``.

Scope: each station's *latest* chiffrier only (the most recent three-year period),
not the full archive back to 1998. The archive is hundreds of 3 MB workbooks;
current conditions are what the chemistry slice is for.

Cache: the package metadata and station list for 30 days; each chiffrier for 90
days (a new period publishes under a new URL, found when the station list
refreshes).
"""

import csv
import io
import json
import logging
import math
import re
import time
import zipfile
from datetime import date, datetime
from pathlib import Path

import httpx
import openpyxl

from src.ingest.discovery import check_resource_discovery
from src.models.water_quality_reading import WaterQualityReading

_PACKAGE_URL = "https://www.donneesquebec.ca/recherche/api/3/action/package_show"
_PACKAGE_ID = "suivi-physicochimique-des-rivieres-et-du-fleuve"
_CACHE_DIR = Path("data/cache/qc_iqbp")
_PACKAGE_TTL = 30 * 86400
_STATIONS_TTL = 30 * 86400
_CHIFFRIER_TTL = 90 * 86400
_USER_AGENT = "fishbot/1.0 (personal fishing exploration bot)"
_JURISDICTION = "CA-QC"
_EARTH_KM_PER_DEGREE = 111.0

_STATIONS_MEMBER = "stations_p.csv"
_SAMPLE_SHEET = "données transposées"

# Column header (casefolded, up to the first space) in the samples sheet -> our field.
_SAMPLE_COLUMNS = {
    "ph": "ph",
    "temp": "temp_c",
    "cond": "conductivity_us_cm",
}

logger = logging.getLogger(__name__)


def fetch_water_quality_readings(
    lat: float,
    lng: float,
    radius_km: float = 50.0,
) -> list[WaterQualityReading]:
    """Fetch Réseau-rivières samples for stations within radius_km of lat/lng."""
    stations = _stations_near(_load_stations(), lat, lng, radius_km)
    if not stations:
        logger.info(
            "QC water quality: no Réseau-rivières stations within %.0fkm of (%.4f, %.4f)",
            radius_km,
            lat,
            lng,
        )
        return []

    by_url: dict[str, dict[str, dict]] = {}
    for station in stations.values():
        by_url.setdefault(station["chiffrier_url"], {})[station["id"]] = station

    readings: list[WaterQualityReading] = []
    failed: list[str] = []
    for url, wanted in by_url.items():
        try:
            path = _download_chiffrier(url)
            readings.extend(parse_chiffrier(path, wanted))
        except Exception as exc:  # noqa: BLE001 - one bad workbook must not drop the rest
            failed.append(url)
            logger.warning("QC water quality: %s failed — %s", url, exc)

    if failed:
        logger.warning(
            "QC water quality: %d of %d chiffriers failed; their stations have no readings",
            len(failed),
            len(by_url),
        )
    logger.info(
        "QC water quality: %d readings from %d stations in %d chiffriers",
        len(readings),
        len(stations),
        len(by_url),
    )
    return readings


# ── station list ───────────────────────────────────────────────────────────────


def _load_stations() -> list[dict]:
    """Parse stations_p.csv into one dict per station (latest year's row)."""
    csv_url = _find_csv_zip_url()
    if not csv_url:
        return []
    zip_bytes = _cached_download(csv_url, _CACHE_DIR / "iqbp_csv.zip", _STATIONS_TTL)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        member = next((n for n in zf.namelist() if n.casefold().endswith(_STATIONS_MEMBER)), None)
        if member is None:
            logger.error(
                "QC water quality: %s not in the CSV zip (members: %s) — the publisher "
                "has changed the archive layout; this is an adapter bug, not an empty source",
                _STATIONS_MEMBER,
                zf.namelist(),
            )
            return []
        text = zf.read(member).decode("utf-8-sig")
    return parse_stations(text)


def parse_stations(text: str) -> list[dict]:
    """Latest-year row per NO_BQMA, with coordinates and the chiffrier URL."""
    latest: dict[str, tuple[int, dict]] = {}
    skipped = 0
    for row in csv.DictReader(io.StringIO(text), delimiter=";"):
        sid = (row.get("NO_BQMA") or "").strip()
        try:
            year = int(row["ANNEE"])
            lat = float(row["LATITUDE"])
            lng = float(row["LONGITUDE"])
        except (KeyError, TypeError, ValueError):
            skipped += 1
            continue
        url = (row.get("URL_ZGIEBV") or row.get("URL_ZGIESL") or "").strip()
        if not sid or not url:
            skipped += 1
            continue
        if sid not in latest or year > latest[sid][0]:
            latest[sid] = (
                year,
                {
                    "id": sid,
                    "name": (row.get("DESCRIPTION") or "").strip() or None,
                    "lat": lat,
                    "lng": lng,
                    "chiffrier_url": url,
                },
            )
    if skipped:
        logger.warning(
            "QC water quality: %d station row(s) skipped (missing id, year, coordinates or "
            "chiffrier URL), %d stations kept",
            skipped,
            len(latest),
        )
    return [station for _, station in latest.values()]


def _stations_near(stations: list[dict], lat: float, lng: float, radius_km: float) -> dict:
    out = {}
    for s in stations:
        if _distance_km(lat, lng, s["lat"], s["lng"]) <= radius_km:
            out[s["id"]] = s
    return out


def _distance_km(lat1: float, lng1: float, lat2: float, lng2: float) -> float:
    dx = (lng2 - lng1) * math.cos(math.radians((lat1 + lat2) / 2))
    return _EARTH_KM_PER_DEGREE * math.hypot(lat2 - lat1, dx)


def _find_csv_zip_url() -> str | None:
    """Locate the CSV zip resource by structure (format + .zip), not by its label."""
    pkg = _cached_json(_PACKAGE_URL, {"id": _PACKAGE_ID}, _CACHE_DIR / "package_meta.json")
    resources = pkg.get("result", {}).get("resources", [])
    matches = [
        r
        for r in resources
        if (r.get("format") or "").strip().upper() == "CSV"
        and (r.get("url") or "").lower().endswith(".zip")
    ]
    check_resource_discovery(
        source="QC Réseau-rivières",
        matched=len(matches),
        candidates=[r.get("name", "") for r in resources],
        matcher="format CSV with a .zip URL",
    )
    return matches[0]["url"] if matches else None


# ── chiffrier parsing ──────────────────────────────────────────────────────────


def parse_chiffrier(path: Path, wanted: dict[str, dict]) -> list[WaterQualityReading]:
    """Read the samples sheet of one chiffrier, keeping only the wanted stations."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next((ws for ws in wb if ws.title.strip().casefold() == _SAMPLE_SHEET), None)
        if sheet is None:
            logger.error(
                "QC water quality: %s has no %r sheet (sheets: %s) — layout changed; "
                "this is an adapter bug, not an empty source",
                path.name,
                _SAMPLE_SHEET,
                wb.sheetnames,
            )
            return []
        sheet.reset_dimensions()  # read-only mode trusts a dimension tag these files get wrong
        rows = sheet.iter_rows(values_only=True)
        columns = _column_map(next(rows, ()))
        if columns is None:
            logger.error(
                "QC water quality: %s samples sheet lacks station/date columns — layout changed",
                path.name,
            )
            return []

        samples: list[WaterQualityReading] = []
        rejected = 0
        for row in rows:
            sample = _row_to_reading(row, columns, wanted)
            if sample is _REJECTED:
                rejected += 1
            elif sample is not None:
                samples.append(sample)
    finally:
        wb.close()

    if rejected:
        logger.warning(
            "QC water quality %s: %d sample(s) rejected by validation (out-of-range values "
            "in the published data), %d kept",
            path.name,
            rejected,
            len(samples),
        )
    return samples


_REJECTED = object()


def _column_map(header: tuple) -> dict[str, int] | None:
    """Map our fields (plus station/date/lab id) to column positions by header text."""
    cols: dict[str, int] = {}
    for i, cell in enumerate(header):
        label = str(cell or "").strip().casefold()
        if not label:
            continue
        key = label.split(" ")[0]
        if label.startswith("n° station"):
            cols["station"] = i
        elif label == "date":
            cols["date"] = i
        elif label.startswith("n° labo"):
            cols["labo"] = i
        elif key in _SAMPLE_COLUMNS:
            cols.setdefault(_SAMPLE_COLUMNS[key], i)
    return cols if {"station", "date"} <= cols.keys() else None


def _row_to_reading(row: tuple, columns: dict[str, int], wanted: dict[str, dict]):
    def cell(name: str):
        i = columns.get(name)
        return row[i] if i is not None and i < len(row) else None

    station_id = str(cell("station") or "").strip()
    station = wanted.get(station_id)
    if station is None:
        return None
    sampled_at = _as_date(cell("date"))
    if sampled_at is None:
        return None

    values = {f: _as_float(cell(f)) for f in ("ph", "temp_c", "conductivity_us_cm")}
    if all(v is None for v in values.values()):
        return None  # a sample of only nutrients/bacteria says nothing about habitability

    labo = str(cell("labo") or "").strip()
    record_id = f"qc_{labo}" if labo else f"qc_{station_id}_{sampled_at.isoformat()}"
    try:
        return WaterQualityReading(
            record_id=record_id,
            station_id=station_id,
            station_name=station["name"],
            lat=station["lat"],
            lng=station["lng"],
            jurisdiction=_JURISDICTION,
            sampled_at=sampled_at,
            **values,
        )
    except ValueError:
        return _REJECTED


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.match(r"\d{4}-\d{2}-\d{2}", value.strip()):
        return date.fromisoformat(value.strip()[:10])
    return None


def _as_float(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None  # "<0.2"-style non-detects are strings and carry no usable number


# ── HTTP + file cache ──────────────────────────────────────────────────────────


def _download_chiffrier(url: str) -> Path:
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", url.split("/IQBP/", 1)[-1])
    dest = _CACHE_DIR / "chiffriers" / name
    _cached_download(url, dest, _CHIFFRIER_TTL)
    return dest


def _cached_download(url: str, dest: Path, ttl: int, params: dict | None = None) -> bytes:
    """Return the body of ``url``, from ``dest`` while it is younger than ``ttl``."""
    if dest.exists() and time.time() - dest.stat().st_mtime < ttl:
        return dest.read_bytes()
    dest.parent.mkdir(parents=True, exist_ok=True)
    r = httpx.get(
        url,
        params=params,
        follow_redirects=True,
        headers={"User-Agent": _USER_AGENT},
        timeout=120,
    )
    r.raise_for_status()
    dest.write_bytes(r.content)
    return r.content


def _cached_json(url: str, params: dict, dest: Path) -> dict:
    return json.loads(_cached_download(url, dest, _PACKAGE_TTL, params))
